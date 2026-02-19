"""
Batch Universe Screener

Runs the analyst agent across all Xover S44 constituents and outputs
results to a sorted Excel workbook with a summary sheet.

Usage:
    python -m scripts.run_universe                    # Full 75-name universe
    python -m scripts.run_universe --limit 5          # Test with 5 names
    python -m scripts.run_universe --index main       # Run Main index instead
"""

import argparse
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from agents.analyst import assess_credit
from core.models import CreditAssessment

INDEX_FILES = {
    "xover": "indices/xover_s44.json",
    "main": "indices/main_s44.json",
}


def load_constituents(index: str) -> list[str]:
    """Load constituent names from index JSON file."""
    path = INDEX_FILES.get(index.lower())
    if not path or not os.path.exists(path):
        print(f"Error: Index file not found for '{index}'", file=sys.stderr)
        sys.exit(1)

    with open(path) as f:
        data = json.load(f)

    names = []
    for sector_names in data["sectors"].values():
        names.extend(sector_names)
    return sorted(names)


def run_screen(
    index: str, limit: int | None = None, delay: float = 2.0
) -> list[CreditAssessment]:
    """Run analyst on each constituent, returning list of assessments."""
    constituents = load_constituents(index)
    if limit:
        constituents = constituents[:limit]

    index_label = "Main" if index.lower() == "main" else "Xover"
    total = len(constituents)
    assessments: list[CreditAssessment] = []
    failures: list[tuple[str, str]] = []

    print(f"Screening {total} names in iTraxx {index_label}\n")

    for i, name in enumerate(constituents, 1):
        print(f"[{i}/{total}] {name}...", end=" ", flush=True)
        try:
            assessment = assess_credit(name, index_label)
            assessments.append(assessment)
            print(f"{assessment.direction.value} (conviction {assessment.conviction})")
        except Exception as e:
            error_msg = str(e)[:80]
            failures.append((name, error_msg))
            print(f"FAILED — {error_msg}")

        if i < total:
            time.sleep(delay)

    print(f"\nCompleted: {len(assessments)}/{total} succeeded, {len(failures)} failed")
    if failures:
        print("Failures:")
        for name, err in failures:
            print(f"  - {name}: {err}")

    return assessments


def force_rank_convictions(assessments: list[CreditAssessment]) -> list[CreditAssessment]:
    """Force-rank conviction by relative mispricing magnitude.

    Instead of relying on Claude's subjective conviction (which clusters at 4/5),
    rank all names by |fair_spread - current_spread| / current_spread and assign
    conviction based on percentile buckets:
        Top 5%   -> 5  (3-4 names out of 75)
        Next 10% -> 4  (7-8 names)
        Mid 50%  -> 3  (35-38 names)
        Next 25% -> 2  (18-19 names)
        Bot 10%  -> 1  (7-8 names)

    Claude's original conviction is preserved in raw_conviction.
    """
    if not assessments:
        return assessments

    n = len(assessments)

    # Compute relative mispricing for each assessment
    ranked = []
    for a in assessments:
        if a.current_spread > 0:
            rel_mispricing = abs(a.fair_spread - a.current_spread) / a.current_spread
        else:
            rel_mispricing = abs(a.fair_spread - a.current_spread) / max(a.fair_spread, 1.0)
        ranked.append((a, rel_mispricing))

    # Sort descending by relative mispricing (biggest edge first)
    ranked.sort(key=lambda x: x[1], reverse=True)

    # Percentile cutoffs
    cut_5 = max(1, round(n * 0.05))         # top 5%   -> conv 5
    cut_15 = max(cut_5 + 1, round(n * 0.15))  # next 10% -> conv 4
    cut_65 = max(cut_15 + 1, round(n * 0.65))  # mid 50%  -> conv 3
    cut_90 = max(cut_65 + 1, round(n * 0.90))  # next 25% -> conv 2
    # bottom 10% -> conv 1

    print(f"\n  Force-ranking convictions across {n} names:")
    print(f"    Conv 5: ranks 1-{cut_5}  ({cut_5} names)")
    print(f"    Conv 4: ranks {cut_5+1}-{cut_15}  ({cut_15 - cut_5} names)")
    print(f"    Conv 3: ranks {cut_15+1}-{cut_65}  ({cut_65 - cut_15} names)")
    print(f"    Conv 2: ranks {cut_65+1}-{cut_90}  ({cut_90 - cut_65} names)")
    print(f"    Conv 1: ranks {cut_90+1}-{n}  ({n - cut_90} names)")

    for idx, (a, rel_misp) in enumerate(ranked):
        # Save Claude's original conviction
        a.raw_conviction = a.conviction

        # Assign force-ranked conviction
        if idx < cut_5:
            a.conviction = 5
        elif idx < cut_15:
            a.conviction = 4
        elif idx < cut_65:
            a.conviction = 3
        elif idx < cut_90:
            a.conviction = 2
        else:
            a.conviction = 1

    # Distribution summary
    dist = {}
    for a, _ in ranked:
        dist[a.conviction] = dist.get(a.conviction, 0) + 1
    print(f"    Distribution: " + ", ".join(f"{k}/5={v}" for k, v in sorted(dist.items(), reverse=True)))

    # Return as flat list (original order doesn't matter, sort_assessments handles final order)
    return [a for a, _ in ranked]


def sort_assessments(assessments: list[CreditAssessment]) -> list[CreditAssessment]:
    """Sort by conviction desc, then absolute mispricing desc."""
    return sorted(
        assessments,
        key=lambda a: (a.conviction, abs(a.fair_spread - a.current_spread)),
        reverse=True,
    )


def write_excel(
    assessments: list[CreditAssessment], index: str, output_dir: str = "outputs"
) -> str:
    """Write assessments to Excel workbook with details + summary sheets."""
    Path(output_dir).mkdir(exist_ok=True)
    index_label = index.lower()
    today = datetime.now().strftime("%Y%m%d")
    filename = f"{index_label}_screen_{today}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb = Workbook()

    # --- Details sheet ---
    ws = wb.active
    ws.title = "Screen Results"

    headers = [
        "Entity Name",
        "Direction",
        "Conviction",
        "Raw Conviction",
        "Current Spread",
        "Fair Spread",
        "Spread Mispricing",
        "Rel Mispricing %",
        "Thesis",
        "Catalyst",
        "Key Risks",
        "Signal Sources",
        "Updated At",
    ]

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    thin_border = Border(
        bottom=Side(style="thin", color="CCCCCC"),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Direction color fills
    long_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    short_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
    flat_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

    sorted_assessments = sort_assessments(assessments)

    for row_idx, a in enumerate(sorted_assessments, 2):
        mispricing = a.fair_spread - a.current_spread
        if a.current_spread > 0:
            rel_misp = abs(mispricing) / a.current_spread * 100
        else:
            rel_misp = abs(mispricing) / max(a.fair_spread, 1.0) * 100

        ws.cell(row=row_idx, column=1, value=a.entity_name)
        dir_cell = ws.cell(row=row_idx, column=2, value=a.direction.value)
        ws.cell(row=row_idx, column=3, value=a.conviction)
        ws.cell(row=row_idx, column=4, value=a.raw_conviction if a.raw_conviction else a.conviction)
        ws.cell(row=row_idx, column=5, value=round(a.current_spread, 1))
        ws.cell(row=row_idx, column=6, value=round(a.fair_spread, 1))
        ws.cell(row=row_idx, column=7, value=round(mispricing, 1))
        ws.cell(row=row_idx, column=8, value=round(rel_misp, 1))
        ws.cell(row=row_idx, column=9, value=a.thesis)
        ws.cell(row=row_idx, column=10, value=a.catalyst)
        ws.cell(row=row_idx, column=11, value="; ".join(a.key_risks))
        ws.cell(row=row_idx, column=12, value=", ".join(a.signal_sources))
        ws.cell(row=row_idx, column=13, value=a.updated_at.strftime("%Y-%m-%d %H:%M"))

        # Color the direction cell
        if a.direction.value == "LONG_RISK":
            dir_cell.fill = long_fill
        elif a.direction.value == "SHORT_RISK":
            dir_cell.fill = short_fill
        else:
            dir_cell.fill = flat_fill

        # Light border on every row
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).border = thin_border

    # Column widths
    col_widths = [35, 14, 12, 14, 14, 12, 16, 16, 60, 50, 60, 40, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Wrap text on thesis/catalyst/risks columns
    for row_idx in range(2, len(sorted_assessments) + 2):
        for col in [9, 10, 11]:
            ws.cell(row=row_idx, column=col).alignment = Alignment(wrap_text=True, vertical="top")

    # Freeze header row
    ws.freeze_panes = "A2"

    # --- Summary sheet ---
    ss = wb.create_sheet("Summary")

    longs = [a for a in sorted_assessments if a.direction.value == "LONG_RISK"]
    shorts = [a for a in sorted_assessments if a.direction.value == "SHORT_RISK"]
    flats = [a for a in sorted_assessments if a.direction.value == "FLAT"]
    avg_conviction = (
        sum(a.conviction for a in sorted_assessments) / len(sorted_assessments)
        if sorted_assessments
        else 0
    )

    title_font = Font(bold=True, size=14, color="1F4E79")
    section_font = Font(bold=True, size=12, color="1F4E79")
    label_font = Font(bold=True, size=11)

    ss.cell(row=1, column=1, value=f"iTraxx {index.capitalize()} Screen — {today}").font = title_font
    ss.merge_cells("A1:D1")

    row = 3
    ss.cell(row=row, column=1, value="Overview").font = section_font
    row += 1
    stats = [
        ("Total Names Screened", len(sorted_assessments)),
        ("Long Risk", len(longs)),
        ("Short Risk", len(shorts)),
        ("Flat", len(flats)),
        ("Average Conviction", round(avg_conviction, 2)),
    ]
    for label, value in stats:
        ss.cell(row=row, column=1, value=label).font = label_font
        ss.cell(row=row, column=2, value=value)
        row += 1

    # Top 5 longs: highest conviction, then most negative mispricing (fair < current = cheap)
    row += 1
    ss.cell(row=row, column=1, value="Top 5 Longs (Best Opportunities)").font = section_font
    row += 1
    top_headers = ["Entity", "Conviction", "Current Spread", "Fair Spread", "Mispricing"]
    for col, h in enumerate(top_headers, 1):
        ss.cell(row=row, column=col, value=h).font = label_font
    row += 1

    top_longs = sorted(
        longs,
        key=lambda a: (a.conviction, -(a.fair_spread - a.current_spread)),
        reverse=True,
    )[:5]
    for a in top_longs:
        ss.cell(row=row, column=1, value=a.entity_name)
        ss.cell(row=row, column=2, value=a.conviction)
        ss.cell(row=row, column=3, value=round(a.current_spread, 1))
        ss.cell(row=row, column=4, value=round(a.fair_spread, 1))
        ss.cell(row=row, column=5, value=round(a.fair_spread - a.current_spread, 1))
        row += 1

    # Top 5 shorts: highest conviction, then most positive mispricing (fair > current = rich)
    row += 1
    ss.cell(row=row, column=1, value="Top 5 Shorts (Highest Risk)").font = section_font
    row += 1
    for col, h in enumerate(top_headers, 1):
        ss.cell(row=row, column=col, value=h).font = label_font
    row += 1

    top_shorts = sorted(
        shorts,
        key=lambda a: (a.conviction, a.fair_spread - a.current_spread),
        reverse=True,
    )[:5]
    for a in top_shorts:
        ss.cell(row=row, column=1, value=a.entity_name)
        ss.cell(row=row, column=2, value=a.conviction)
        ss.cell(row=row, column=3, value=round(a.current_spread, 1))
        ss.cell(row=row, column=4, value=round(a.fair_spread, 1))
        ss.cell(row=row, column=5, value=round(a.fair_spread - a.current_spread, 1))
        row += 1

    # Summary column widths
    ss.column_dimensions["A"].width = 40
    ss.column_dimensions["B"].width = 14
    ss.column_dimensions["C"].width = 16
    ss.column_dimensions["D"].width = 14
    ss.column_dimensions["E"].width = 14

    wb.save(filepath)
    return filepath


def main():
    parser = argparse.ArgumentParser(description="Batch Universe Screener")
    parser.add_argument(
        "--index",
        type=str,
        choices=["xover", "main"],
        default="xover",
        help="Index to screen (default: xover)",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Limit to first N names (for testing)",
    )
    parser.add_argument(
        "--delay",
        type=float,
        default=2.0,
        help="Seconds between API calls (default: 2.0)",
    )
    args = parser.parse_args()

    assessments = run_screen(args.index, limit=args.limit, delay=args.delay)

    if not assessments:
        print("No assessments produced. Exiting.", file=sys.stderr)
        sys.exit(1)

    # Force-rank convictions by relative mispricing instead of Claude's subjective scoring
    assessments = force_rank_convictions(assessments)

    filepath = write_excel(assessments, args.index)
    print(f"\nExcel saved: {filepath}")


if __name__ == "__main__":
    main()
