#!/usr/bin/env python3
"""
Import iTraxx S44 constituent spreads from Bloomberg Excel extract.

Usage:
    python scripts/import_spreads.py "path/to/ITRX S44.xlsx"

Reads MAIN S44 and XOVER S44 sheets, merges equity ticker data from
data/itraxx_equity_map.json, and writes data/itraxx_s44_master.json.
"""

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
EQUITY_MAP_PATH = PROJECT_ROOT / "data" / "itraxx_equity_map.json"
OUTPUT_PATH = PROJECT_ROOT / "data" / "itraxx_s44_master.json"


def parse_sheet(headers: list, rows: list, index_name: str) -> dict:
    constituents = {}
    has_upfront = len(headers) >= 9

    for row in rows:
        if not row or not row[0]:
            continue
        name = str(row[0]).strip()
        convention = str(row[6] or "Spread").strip().lower().replace(" ", "_")
        spread = row[7] if row[7] is not None else None
        upfront = row[8] if has_upfront and len(row) > 8 else None

        if spread is not None:
            try:
                spread = round(float(spread), 3)
            except (ValueError, TypeError):
                spread = None

        if upfront is not None:
            try:
                upfront = round(float(upfront), 2)
            except (ValueError, TypeError):
                upfront = None

        constituents[name] = {
            "index": index_name,
            "weight": round(float(row[1]), 4) if row[1] else None,
            "corp_ticker": str(row[2]).strip() if row[2] else None,
            "cds_ticker": str(row[3]).strip() if row[3] else None,
            "isin": str(row[4]).strip() if row[4] else None,
            "red_pair": str(row[5]).strip() if row[5] else None,
            "quote_convention": convention,
            "spread_bps": spread,
            "points_upfront": upfront if convention == "points_upfront" else None,
            "equity_ticker": None,
            "is_public": None,
            "options_available": None,
            "options_liquidity": None,
            "sector": None,
            "exchange": None,
            "currency": None,
        }

    return constituents


def _normalise(name: str) -> str:
    return (
        name.lower()
        .replace("/", " ")
        .replace("-", " ")
        .replace("  ", " ")
        .strip()
    )


def merge_equity_data(constituents: dict, equity_map: dict) -> None:
    norm_lookup = {_normalise(k): v for k, v in equity_map.items()}

    equity_fields = [
        "equity_ticker", "is_public", "options_available",
        "options_liquidity", "sector", "exchange", "currency",
    ]

    for name, data in constituents.items():
        eq = equity_map.get(name)
        if not eq:
            eq = norm_lookup.get(_normalise(name))
        if not eq:
            # Fuzzy: try prefix / substring matching on normalised names
            norm_name = _normalise(name)
            for norm_key, val in norm_lookup.items():
                if norm_name.startswith(norm_key) or norm_key.startswith(norm_name):
                    eq = val
                    break
        if eq:
            for field in equity_fields:
                if field in eq and eq[field] is not None:
                    data[field] = eq[field]


def compute_metadata(constituents: dict, snapshot_date: str = None) -> dict:
    main = [c for c in constituents.values() if c["index"] == "main"]
    xover = [c for c in constituents.values() if c["index"] == "xover"]

    main_spreads = [c["spread_bps"] for c in main if c["spread_bps"] is not None]
    xover_spreads = [c["spread_bps"] for c in xover if c["spread_bps"] is not None]

    return {
        "series": 44,
        "snapshot_date": snapshot_date or str(date.today()),
        "source": "Bloomberg ITRX S44 Constituents",
        "main_count": len(main),
        "xover_count": len(xover),
        "main_avg_spread": round(sum(main_spreads) / len(main_spreads), 2) if main_spreads else 0,
        "xover_avg_spread": round(sum(xover_spreads) / len(xover_spreads), 2) if xover_spreads else 0,
        "main_median_spread": round(sorted(main_spreads)[len(main_spreads) // 2], 2) if main_spreads else 0,
        "xover_median_spread": round(sorted(xover_spreads)[len(xover_spreads) // 2], 2) if xover_spreads else 0,
    }


def import_from_excel(excel_path: str, snapshot_date: str = None) -> dict:
    wb = load_workbook(excel_path, read_only=True, data_only=True)

    all_constituents = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        headers = list(next(rows_iter))
        data_rows = [list(r) for r in rows_iter]

        if "MAIN" in sheet_name.upper():
            idx = "main"
        elif "XOVER" in sheet_name.upper():
            idx = "xover"
        else:
            continue

        parsed = parse_sheet(headers, data_rows, index_name=idx)
        all_constituents.update(parsed)

    wb.close()

    if EQUITY_MAP_PATH.exists():
        with open(EQUITY_MAP_PATH) as f:
            eq_data = json.load(f)
        equity_names = eq_data.get("names", eq_data)
        merge_equity_data(all_constituents, equity_names)

    metadata = compute_metadata(all_constituents, snapshot_date)

    return {"metadata": metadata, "constituents": all_constituents}


def main():
    if len(sys.argv) < 2:
        print("Usage: python scripts/import_spreads.py <path-to-excel>")
        sys.exit(1)

    excel_path = sys.argv[1]
    if not Path(excel_path).exists():
        print(f"File not found: {excel_path}")
        sys.exit(1)

    filename = Path(excel_path).stem
    snapshot_date = None
    if filename[:4].isdigit() and len(filename) >= 10:
        snapshot_date = filename[:10]

    master = import_from_excel(excel_path, snapshot_date)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w") as f:
        json.dump(master, f, indent=2, ensure_ascii=False)

    meta = master["metadata"]
    print(f"\n{'=' * 60}")
    print(f"  iTraxx S44 Master Import — {meta['snapshot_date']}")
    print(f"{'=' * 60}")
    print(f"  Main:  {meta['main_count']} names  avg {meta['main_avg_spread']:.1f}bp  med {meta['main_median_spread']:.1f}bp")
    print(f"  Xover: {meta['xover_count']} names  avg {meta['xover_avg_spread']:.1f}bp  med {meta['xover_median_spread']:.1f}bp")
    print(f"  Total: {meta['main_count'] + meta['xover_count']} constituents")

    distressed = [
        (n, c) for n, c in master["constituents"].items()
        if c.get("spread_bps") and c["spread_bps"] > 500
    ]
    if distressed:
        print(f"\n  Distressed (>500bp):")
        for name, c in sorted(distressed, key=lambda x: -x[1]["spread_bps"]):
            uf = f"  ({c['points_upfront']}pts)" if c.get("points_upfront") else ""
            print(f"    {name:<40} {c['spread_bps']:>8.1f}bp{uf}")

    print(f"\n  Written to: {OUTPUT_PATH}")
    print(f"{'=' * 60}\n")


if __name__ == "__main__":
    main()
