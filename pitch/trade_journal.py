"""
Trade Journal — SQLite-backed trade logging and tracking.

Logs every recommendation from the universe screen as an "open" trade,
tracks P&L, and supports post-mortems for closed positions.

Database: data/trade_journal.db

Usage:
    python -m pitch.trade_journal --init           # Load today's screen as open trades
    python -m pitch.trade_journal --status          # Show all open trades
    python -m pitch.trade_journal --all             # Show all trades (open + closed)
    python -m pitch.trade_journal --close TRADE_ID --exit-spread 350 --post-mortem "Thesis played out"
    python -m pitch.trade_journal --stats           # Portfolio-level stats
"""

import argparse
import os
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

DB_PATH = Path("data/trade_journal.db")

# ---------------------------------------------------------------------------
# Schema
# ---------------------------------------------------------------------------

CREATE_TABLE = """
CREATE TABLE IF NOT EXISTS trades (
    trade_id        INTEGER PRIMARY KEY AUTOINCREMENT,
    entity_name     TEXT NOT NULL,
    direction       TEXT NOT NULL,
    conviction      INTEGER NOT NULL,
    raw_conviction  INTEGER,
    entry_spread    REAL NOT NULL,
    fair_spread     REAL,
    mispricing_bps  REAL,
    rel_mispricing  REAL,
    thesis          TEXT,
    catalyst        TEXT,
    key_risks       TEXT,
    signal_sources  TEXT,
    strategy_type   TEXT DEFAULT 'single_name_cds',
    signal_source   TEXT DEFAULT 'universe_screen',
    entry_date      TEXT NOT NULL,
    exit_date       TEXT,
    exit_spread     REAL,
    pnl_bps         REAL,
    status          TEXT NOT NULL DEFAULT 'open',
    post_mortem     TEXT,
    created_at      TEXT NOT NULL DEFAULT (datetime('now')),
    updated_at      TEXT NOT NULL DEFAULT (datetime('now'))
);
"""

CREATE_INDEX = """
CREATE INDEX IF NOT EXISTS idx_trades_status ON trades(status);
CREATE INDEX IF NOT EXISTS idx_trades_entity ON trades(entity_name);
CREATE INDEX IF NOT EXISTS idx_trades_entry_date ON trades(entry_date);
"""


# ---------------------------------------------------------------------------
# Database helpers
# ---------------------------------------------------------------------------

def get_connection() -> sqlite3.Connection:
    """Get SQLite connection, creating DB and tables if needed."""
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.executescript(CREATE_TABLE)
    conn.executescript(CREATE_INDEX)
    return conn


def find_latest_screen() -> str | None:
    """Find the most recent xover_screen Excel file."""
    outputs = Path("outputs")
    if not outputs.exists():
        return None
    candidates = sorted(
        outputs.glob("xover_screen_*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    return str(candidates[0]) if candidates else None


# ---------------------------------------------------------------------------
# Init: Load screen into journal
# ---------------------------------------------------------------------------

def init_trades(screen_path: str | None = None) -> int:
    """Load all recommendations from the latest universe screen as open trades.

    Skips names that already have an open trade for the same entry date.

    Returns:
        Number of new trades inserted.
    """
    if not screen_path:
        screen_path = find_latest_screen()
    if not screen_path or not os.path.exists(screen_path):
        print("Error: No screen file found. Run scripts/run_universe.py first.",
              file=sys.stderr)
        return 0

    wb = load_workbook(screen_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    today = datetime.now().strftime("%Y-%m-%d")
    conn = get_connection()

    # Check existing open trades for today
    existing = set()
    for row in conn.execute(
        "SELECT entity_name FROM trades WHERE entry_date = ? AND status = 'open'",
        (today,),
    ):
        existing.add(row["entity_name"])

    inserted = 0
    for row in rows:
        if not row or not row[0]:
            continue

        entity_name = str(row[0]).strip()
        if entity_name in existing:
            continue

        direction = str(row[1] or "").strip()
        conviction = int(row[2]) if row[2] is not None else 3
        raw_conviction = int(row[3]) if row[3] is not None else None
        current_spread = float(row[4]) if row[4] is not None else 0.0
        fair_spread = float(row[5]) if row[5] is not None else 0.0
        mispricing = float(row[6]) if row[6] is not None else 0.0
        rel_mispricing = float(row[7]) if row[7] is not None else 0.0
        thesis = str(row[8] or "").strip()
        catalyst = str(row[9] or "").strip()
        key_risks = str(row[10] or "").strip()
        signal_sources = str(row[11] or "").strip()

        # Determine strategy type from direction
        if direction == "FLAT":
            strategy_type = "monitor"
        else:
            strategy_type = "single_name_cds"

        conn.execute("""
            INSERT INTO trades (
                entity_name, direction, conviction, raw_conviction,
                entry_spread, fair_spread, mispricing_bps, rel_mispricing,
                thesis, catalyst, key_risks, signal_sources,
                strategy_type, signal_source, entry_date, status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            entity_name, direction, conviction, raw_conviction,
            current_spread, fair_spread, mispricing, rel_mispricing,
            thesis, catalyst, key_risks, signal_sources,
            strategy_type, "universe_screen", today, "open",
        ))
        inserted += 1

    conn.commit()
    conn.close()
    return inserted


# ---------------------------------------------------------------------------
# Status: Show trades
# ---------------------------------------------------------------------------

def show_trades(status_filter: str | None = "open", entity_filter: str | None = None):
    """Print trades to terminal."""
    conn = get_connection()

    query = "SELECT * FROM trades"
    params = []
    conditions = []

    if status_filter:
        conditions.append("status = ?")
        params.append(status_filter)

    if entity_filter:
        conditions.append("entity_name LIKE ?")
        params.append(f"%{entity_filter}%")

    if conditions:
        query += " WHERE " + " AND ".join(conditions)

    query += " ORDER BY conviction DESC, entry_spread DESC"
    rows = conn.execute(query, params).fetchall()
    conn.close()

    if not rows:
        print(f"No trades found (filter: status={status_filter or 'all'})")
        return

    # Count by direction
    longs = [r for r in rows if r["direction"] == "LONG_RISK"]
    shorts = [r for r in rows if r["direction"] == "SHORT_RISK"]
    flats = [r for r in rows if r["direction"] == "FLAT"]

    status_label = status_filter.upper() if status_filter else "ALL"
    print(f"\n{'='*100}")
    print(f"  TRADE JOURNAL — {status_label} TRADES ({len(rows)} total | "
          f"{len(longs)} long | {len(shorts)} short | {len(flats)} flat)")
    print(f"{'='*100}")

    print(f"\n  {'ID':>4}  {'Entity':<35} {'Dir':<12} {'Conv':>4} {'Entry':>7} "
          f"{'Fair':>7} {'Misprice':>8} {'Status':<8} {'Date':<12}")
    print(f"  {'-'*4}  {'-'*35} {'-'*12} {'-'*4} {'-'*7} {'-'*7} {'-'*8} {'-'*8} {'-'*12}")

    for r in rows:
        mispricing = r["mispricing_bps"] or 0
        direction_short = {
            "LONG_RISK": "LONG",
            "SHORT_RISK": "SHORT",
            "FLAT": "FLAT",
        }.get(r["direction"], r["direction"][:8])

        status_str = r["status"].upper()

        print(f"  {r['trade_id']:>4}  {r['entity_name'][:35]:<35} "
              f"{direction_short:<12} {r['conviction']:>4} "
              f"{r['entry_spread']:>7.1f} {(r['fair_spread'] or 0):>7.1f} "
              f"{mispricing:>+8.0f} {status_str:<8} {r['entry_date']:<12}")

        # Show exit info for closed trades
        if r["status"] == "closed" and r["exit_spread"]:
            pnl = r["pnl_bps"] or 0
            pnl_color = "+" if pnl > 0 else ""
            print(f"        Exit: {r['exit_spread']:.1f} on {r['exit_date'] or '?'} | "
                  f"P&L: {pnl_color}{pnl:.1f}bps")
            if r["post_mortem"]:
                print(f"        Post-mortem: {r['post_mortem']}")

    print(f"\n{'='*100}")


# ---------------------------------------------------------------------------
# Close trade
# ---------------------------------------------------------------------------

def close_trade(
    trade_id: int,
    exit_spread: float,
    post_mortem: str = "",
    exit_date: str | None = None,
):
    """Close a trade and compute P&L."""
    conn = get_connection()

    # Fetch trade
    trade = conn.execute(
        "SELECT * FROM trades WHERE trade_id = ?", (trade_id,)
    ).fetchone()

    if not trade:
        print(f"Error: Trade ID {trade_id} not found.", file=sys.stderr)
        conn.close()
        return

    if trade["status"] == "closed":
        print(f"Warning: Trade ID {trade_id} is already closed.", file=sys.stderr)
        conn.close()
        return

    # Compute P&L in bps
    entry = trade["entry_spread"]
    direction = trade["direction"]

    if direction == "LONG_RISK":
        # Protection seller: profit when spreads tighten
        pnl_bps = entry - exit_spread
    elif direction == "SHORT_RISK":
        # Protection buyer: profit when spreads widen
        pnl_bps = exit_spread - entry
    else:
        pnl_bps = 0.0

    ex_date = exit_date or datetime.now().strftime("%Y-%m-%d")

    conn.execute("""
        UPDATE trades SET
            status = 'closed',
            exit_spread = ?,
            exit_date = ?,
            pnl_bps = ?,
            post_mortem = ?,
            updated_at = datetime('now')
        WHERE trade_id = ?
    """, (exit_spread, ex_date, pnl_bps, post_mortem, trade_id))

    conn.commit()
    conn.close()

    pnl_sign = "+" if pnl_bps > 0 else ""
    print(f"\nTrade #{trade_id} closed:")
    print(f"  Entity:      {trade['entity_name']}")
    print(f"  Direction:   {trade['direction']}")
    print(f"  Entry:       {entry:.1f}bps on {trade['entry_date']}")
    print(f"  Exit:        {exit_spread:.1f}bps on {ex_date}")
    print(f"  P&L:         {pnl_sign}{pnl_bps:.1f}bps")
    if post_mortem:
        print(f"  Post-mortem: {post_mortem}")


# ---------------------------------------------------------------------------
# Stats
# ---------------------------------------------------------------------------

def show_stats():
    """Show portfolio-level trade statistics."""
    conn = get_connection()

    total = conn.execute("SELECT COUNT(*) as n FROM trades").fetchone()["n"]
    open_count = conn.execute(
        "SELECT COUNT(*) as n FROM trades WHERE status = 'open'"
    ).fetchone()["n"]
    closed_count = conn.execute(
        "SELECT COUNT(*) as n FROM trades WHERE status = 'closed'"
    ).fetchone()["n"]

    # Closed trade stats
    closed = conn.execute(
        "SELECT * FROM trades WHERE status = 'closed' AND pnl_bps IS NOT NULL"
    ).fetchall()

    conn.close()

    print(f"\n{'='*60}")
    print(f"  TRADE JOURNAL — STATISTICS")
    print(f"{'='*60}")
    print(f"  Total trades:      {total}")
    print(f"  Open:              {open_count}")
    print(f"  Closed:            {closed_count}")

    if closed:
        pnls = [r["pnl_bps"] for r in closed]
        winners = [p for p in pnls if p > 0]
        losers = [p for p in pnls if p < 0]

        total_pnl = sum(pnls)
        avg_pnl = total_pnl / len(pnls)
        win_rate = len(winners) / len(pnls) * 100
        avg_win = sum(winners) / len(winners) if winners else 0
        avg_loss = sum(losers) / len(losers) if losers else 0

        print(f"\n  CLOSED TRADE PERFORMANCE:")
        print(f"  Total P&L:         {total_pnl:+.1f}bps")
        print(f"  Avg P&L:           {avg_pnl:+.1f}bps")
        print(f"  Win rate:          {win_rate:.0f}%")
        print(f"  Winners:           {len(winners)} (avg {avg_win:+.1f}bps)")
        print(f"  Losers:            {len(losers)} (avg {avg_loss:+.1f}bps)")
        if avg_loss != 0:
            print(f"  Win/Loss ratio:    {abs(avg_win / avg_loss):.2f}x")
        print(f"  Best:              {max(pnls):+.1f}bps")
        print(f"  Worst:             {min(pnls):+.1f}bps")
    else:
        print(f"\n  No closed trades yet.")

    print(f"{'='*60}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Trade Journal")
    parser.add_argument(
        "--init", action="store_true",
        help="Load today's screen as open trades",
    )
    parser.add_argument(
        "--screen-file", type=str, default=None,
        help="Specific screen Excel file for --init",
    )
    parser.add_argument(
        "--status", action="store_true",
        help="Show all open trades",
    )
    parser.add_argument(
        "--all", action="store_true",
        help="Show all trades (open + closed)",
    )
    parser.add_argument(
        "--entity", type=str, default=None,
        help="Filter by entity name (partial match)",
    )
    parser.add_argument(
        "--close", type=int, default=None,
        help="Close a trade by ID",
    )
    parser.add_argument(
        "--exit-spread", type=float, default=None,
        help="Exit spread for --close",
    )
    parser.add_argument(
        "--post-mortem", type=str, default="",
        help="Post-mortem note for --close",
    )
    parser.add_argument(
        "--stats", action="store_true",
        help="Show trade statistics",
    )
    args = parser.parse_args()

    if args.init:
        print(f"Initialising trade journal from universe screen...")
        screen = args.screen_file or find_latest_screen()
        print(f"Screen file: {screen}")
        count = init_trades(screen)
        print(f"Inserted {count} new trades into {DB_PATH}")
        print(f"\nUse --status to view open trades.")
        return

    if args.close is not None:
        if args.exit_spread is None:
            print("Error: --exit-spread required for --close", file=sys.stderr)
            sys.exit(1)
        close_trade(args.close, args.exit_spread, args.post_mortem)
        return

    if args.stats:
        show_stats()
        return

    if args.all:
        show_trades(status_filter=None, entity_filter=args.entity)
        return

    if args.status or not any([args.init, args.close, args.stats, args.all]):
        show_trades(status_filter="open", entity_filter=args.entity)
        return


if __name__ == "__main__":
    main()
