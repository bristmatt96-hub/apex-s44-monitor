"""
Simulated Track Record Generator — Strategies in Credit Backtest

Takes today's 75 assessments + strategist portfolio and simulates a 3-month
forward P&L track assuming spreads mean-revert 50% toward fair value with
stochastic daily noise.

Outputs:
    outputs/backtest/backtest_report_YYYYMMDD.xlsx  (5 sheets)
    outputs/backtest/backtest_report_YYYYMMDD.pdf   (one-page summary)

Usage:
    python -m pitch.backtest_report                        # Default 63 trading days
    python -m pitch.backtest_report --days 126             # 6-month simulation
    python -m pitch.backtest_report --seed 42              # Reproducible
    python -m pitch.backtest_report --historical           # Placeholder for real data
    python -m pitch.backtest_report --json                 # JSON output only
"""

import argparse
import json
import math
import os
import sys
from dataclasses import dataclass, field, asdict
from datetime import datetime, timedelta
from pathlib import Path
from statistics import mean, median, stdev

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm, mm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from analytics.cds_pricer import spread_to_upfront, _risky_annuity

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

OUTPUT_DIR = Path("outputs/backtest")

# Simulation defaults
DEFAULT_TRADING_DAYS = 63         # ~3 months
DAILY_VOL_BPS = 5.0              # Normal daily noise σ
MEAN_REVERSION_HALF_LIFE = 63    # 50% reversion over horizon
NAV_MILLIONS = 500.0
RECOVERY = 0.40
MATURITY = 5.0
COUPON_BPS = 500.0               # HY standard
RISK_FREE = 0.03

# Rating bucket boundaries (spread-implied)
RATING_THRESHOLDS = {
    "IG":  (0, 150),
    "BB":  (150, 300),
    "B":   (300, 500),
    "B-":  (500, 800),
    "CCC": (800, 5000),
}

# Alpha source keywords (mapped from thesis/signal text)
ALPHA_SOURCES = ["RV", "Catalyst", "Filing Speed"]

# Excel styling
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
SECTION_FONT = Font(bold=True, size=12, color="1F4E79")
LABEL_FONT = Font(bold=True, size=10)
LONG_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
SHORT_FILL = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
GREEN_FONT = Font(bold=True, color="28A745")
RED_FONT = Font(bold=True, color="DC3545")
THIN_BORDER = Border(bottom=Side(style="thin", color="CCCCCC"))
NUMBER_FMT_2DP = "#,##0.00"
NUMBER_FMT_0DP = "#,##0"
PCT_FMT = "0.0%"

# PDF colour palette (matches idea_sheet.py)
DARK_NAVY = colors.HexColor("#0D1B2A")
NAVY = colors.HexColor("#1B2838")
STEEL = colors.HexColor("#415A77")
LIGHT_STEEL = colors.HexColor("#778DA9")
OFF_WHITE = colors.HexColor("#F8F9FA")
LONG_GREEN = colors.HexColor("#28A745")
SHORT_RED = colors.HexColor("#DC3545")
ACCENT = colors.HexColor("#1976D2")
DIVIDER = colors.HexColor("#DEE2E6")


# ---------------------------------------------------------------------------
# Data models
# ---------------------------------------------------------------------------

@dataclass
class SimPosition:
    """A single position in the backtest simulation."""
    entity_name: str
    direction: str           # LONG_RISK or SHORT_RISK
    conviction: int          # 1-5
    notional_m: float        # Millions
    entry_spread: float      # bps
    fair_spread: float       # bps
    sector: str
    estimated_rating: str
    alpha_source: str        # RV, Catalyst, Filing Speed
    thesis: str

    # Simulation state (updated daily)
    current_spread: float = 0.0
    peak_spread: float = 0.0   # For drawdown tracking
    daily_pnl: list = field(default_factory=list)
    spread_path: list = field(default_factory=list)

    def __post_init__(self):
        if self.current_spread == 0.0:
            self.current_spread = self.entry_spread
        self.peak_spread = self.entry_spread


@dataclass
class BacktestMetrics:
    """Overall backtest performance summary."""
    total_return_bps: float = 0.0
    total_return_pct: float = 0.0
    annualised_return_pct: float = 0.0
    max_drawdown_pct: float = 0.0
    max_drawdown_bps: float = 0.0
    sharpe_ratio: float = 0.0
    sortino_ratio: float = 0.0
    calmar_ratio: float = 0.0
    win_rate_pct: float = 0.0
    profit_factor: float = 0.0
    avg_daily_return_bps: float = 0.0
    daily_vol_bps: float = 0.0
    total_pnl_usd: float = 0.0
    best_day_bps: float = 0.0
    worst_day_bps: float = 0.0
    winning_days: int = 0
    losing_days: int = 0
    trading_days: int = 0


@dataclass
class AttributionBucket:
    """P&L attribution for a grouping (sector, rating, alpha source, conviction)."""
    label: str
    total_pnl_usd: float = 0.0
    total_pnl_bps: float = 0.0
    position_count: int = 0
    win_count: int = 0
    loss_count: int = 0
    avg_return_bps: float = 0.0
    contribution_pct: float = 0.0


# ---------------------------------------------------------------------------
# Data loading (reuses patterns from scenario_analysis.py)
# ---------------------------------------------------------------------------

def find_latest_portfolio() -> str | None:
    """Find the most recent portfolio JSON."""
    portfolio_dir = Path("outputs/portfolio")
    if not portfolio_dir.exists():
        return None
    candidates = sorted(
        portfolio_dir.glob("portfolio_*.json"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    return str(candidates[0]) if candidates else None


def load_portfolio(path: str) -> dict:
    """Load portfolio JSON."""
    with open(path) as f:
        return json.load(f)


def load_sector_mapping() -> dict[str, str]:
    """Load entity -> sector from xover_s44.json."""
    path = "indices/xover_s44.json"
    if not os.path.exists(path):
        return {}
    with open(path) as f:
        data = json.load(f)
    mapping = {}
    for sector, names in data.get("sectors", {}).items():
        for name in names:
            mapping[name] = sector
    return mapping


def classify_rating(spread: float) -> str:
    """Classify a name into a rating bucket by its spread."""
    if spread < 150:
        return "IG"
    elif spread < 300:
        return "BB"
    elif spread < 500:
        return "B"
    elif spread < 800:
        return "B-"
    else:
        return "CCC"


def infer_alpha_source(thesis: str, direction: str, fair_spread: float,
                       current_spread: float) -> str:
    """Infer which alpha source drives this trade."""
    thesis_lower = thesis.lower() if thesis else ""

    # Filing / regulatory catalyst keywords
    filing_keywords = ["filing", "annual report", "companies house",
                       "regulatory", "restructuring", "lme", "announcement"]
    if any(k in thesis_lower for k in filing_keywords):
        return "Filing Speed"

    # Event / catalyst keywords
    catalyst_keywords = ["earnings", "results", "rebalancing", "refinancing",
                         "downgrade", "upgrade", "tax", "catalyst", "integration"]
    if any(k in thesis_lower for k in catalyst_keywords):
        return "Catalyst"

    # Default: relative value (mispricing driven)
    return "RV"


def build_positions(portfolio: dict) -> list[SimPosition]:
    """Convert portfolio JSON into SimPosition list."""
    positions = []
    for p in portfolio.get("top_positions", []):
        alpha = infer_alpha_source(
            p.get("thesis", ""),
            p.get("direction", ""),
            p.get("fair_spread", 0),
            p.get("current_spread", 0),
        )
        positions.append(SimPosition(
            entity_name=p["entity_name"],
            direction=p["direction"],
            conviction=p.get("conviction", 3),
            notional_m=p.get("notional_millions", 10.0),
            entry_spread=p["current_spread"],
            fair_spread=p["fair_spread"],
            sector=p.get("sector", "Unknown"),
            estimated_rating=p.get("estimated_rating", classify_rating(p["current_spread"])),
            alpha_source=alpha,
            thesis=p.get("thesis", ""),
            current_spread=p["current_spread"],
        ))
    return positions


# ---------------------------------------------------------------------------
# CDS P&L engine (exact re-pricing, same as scenario_analysis.py)
# ---------------------------------------------------------------------------

def compute_daily_pnl(
    spread_yesterday: float,
    spread_today: float,
    notional: float,
    direction: str,
) -> float:
    """Compute exact P&L by re-pricing the CDS at today's spread.

    P&L = (upfront_today - upfront_yesterday) * notional / 100
    SHORT_RISK (protection buyer) profits on widening.
    LONG_RISK (protection seller) profits on tightening.
    """
    rpv01_y = _risky_annuity(spread_yesterday, RECOVERY, MATURITY)
    rpv01_t = _risky_annuity(spread_today, RECOVERY, MATURITY)

    uf_y = (spread_yesterday - COUPON_BPS) / 10_000 * rpv01_y * 100
    uf_t = (spread_today - COUPON_BPS) / 10_000 * rpv01_t * 100

    delta_pct = uf_t - uf_y
    delta_dollars = delta_pct / 100 * notional

    is_prot_buyer = direction == "SHORT_RISK"
    return delta_dollars if is_prot_buyer else -delta_dollars


# ---------------------------------------------------------------------------
# Monte Carlo spread simulation
# ---------------------------------------------------------------------------

def simulate_spread_paths(
    positions: list[SimPosition],
    n_days: int,
    rng: np.random.Generator,
) -> None:
    """Simulate daily spread paths with mean-reversion + noise.

    Model: dS = κ(θ - S)dt + σ·S^0.5·dW

    where:
        κ   = mean-reversion speed (calibrated for 50% reversion over horizon)
        θ   = fair_spread (target)
        σ   = daily volatility (scaled by sqrt of spread level)
        dW  = standard normal noise
    """
    # Mean-reversion speed: κ such that (1 - κ)^n_days = 0.5
    kappa = 1.0 - 0.5 ** (1.0 / n_days)

    for pos in positions:
        s = pos.entry_spread
        pos.spread_path = [s]
        pos.daily_pnl = []

        for day in range(n_days):
            # Mean-reversion pull
            pull = kappa * (pos.fair_spread - s)

            # Stochastic noise: σ scaled by sqrt(spread/300) to keep vol proportional
            vol_scale = max(0.3, math.sqrt(s / 300.0))
            noise = rng.normal(0, DAILY_VOL_BPS * vol_scale)

            # New spread (floor at 5bp)
            s_new = max(5.0, s + pull + noise)

            # Compute P&L for this day
            pnl = compute_daily_pnl(
                s, s_new,
                pos.notional_m * 1_000_000,
                pos.direction,
            )

            pos.daily_pnl.append(pnl)
            pos.spread_path.append(s_new)
            s = s_new

        pos.current_spread = s


# ---------------------------------------------------------------------------
# Metrics calculation
# ---------------------------------------------------------------------------

def compute_metrics(
    positions: list[SimPosition],
    n_days: int,
) -> BacktestMetrics:
    """Compute portfolio-level backtest metrics from simulated daily P&L."""

    nav = NAV_MILLIONS * 1_000_000  # $500M

    # Aggregate daily P&L across all positions
    daily_pnl = [0.0] * n_days
    for pos in positions:
        for d in range(n_days):
            daily_pnl[d] += pos.daily_pnl[d]

    # Convert to bps of NAV
    daily_bps = [p / nav * 10_000 for p in daily_pnl]

    total_pnl = sum(daily_pnl)
    total_bps = sum(daily_bps)

    # Cumulative P&L for drawdown
    cum = [0.0]
    for p in daily_pnl:
        cum.append(cum[-1] + p)

    peak = cum[0]
    max_dd = 0.0
    for c in cum:
        peak = max(peak, c)
        dd = (peak - c) / nav * 100 if peak > 0 else 0.0
        max_dd = max(max_dd, dd)

    # Win/loss days
    winning = [d for d in daily_bps if d > 0]
    losing = [d for d in daily_bps if d < 0]

    avg_daily = mean(daily_bps) if daily_bps else 0
    vol_daily = stdev(daily_bps) if len(daily_bps) > 1 else 1.0

    # Sharpe: annualised (252 days)
    sharpe = (avg_daily * math.sqrt(252)) / vol_daily if vol_daily > 0 else 0.0

    # Sortino: only downside deviation
    downside = [d for d in daily_bps if d < 0]
    downside_dev = math.sqrt(mean([d**2 for d in downside])) if downside else 1.0
    sortino = (avg_daily * math.sqrt(252)) / downside_dev if downside_dev > 0 else 0.0

    # Calmar
    annualised_pct = total_bps / 10000 * (252 / max(n_days, 1)) * 100
    calmar = annualised_pct / max_dd if max_dd > 0 else 0.0

    # Profit factor
    gross_wins = sum(d for d in daily_bps if d > 0)
    gross_losses = abs(sum(d for d in daily_bps if d < 0))
    pf = gross_wins / gross_losses if gross_losses > 0 else 99.0

    return BacktestMetrics(
        total_return_bps=round(total_bps, 2),
        total_return_pct=round(total_bps / 10000 * 100, 4),
        annualised_return_pct=round(annualised_pct, 2),
        max_drawdown_pct=round(max_dd, 4),
        max_drawdown_bps=round(max_dd * 100, 2),
        sharpe_ratio=round(sharpe, 2),
        sortino_ratio=round(sortino, 2),
        calmar_ratio=round(calmar, 2),
        win_rate_pct=round(len(winning) / max(n_days, 1) * 100, 1),
        profit_factor=round(pf, 2),
        avg_daily_return_bps=round(avg_daily, 2),
        daily_vol_bps=round(vol_daily, 2),
        total_pnl_usd=round(total_pnl, 0),
        best_day_bps=round(max(daily_bps), 2) if daily_bps else 0,
        worst_day_bps=round(min(daily_bps), 2) if daily_bps else 0,
        winning_days=len(winning),
        losing_days=len(losing),
        trading_days=n_days,
    )


def compute_attribution(
    positions: list[SimPosition],
    group_key: str,           # "sector", "estimated_rating", "alpha_source", "conviction"
    total_pnl: float,
) -> list[AttributionBucket]:
    """Compute P&L attribution by grouping key."""

    nav = NAV_MILLIONS * 1_000_000
    buckets: dict[str, AttributionBucket] = {}

    for pos in positions:
        label = str(getattr(pos, group_key, "Unknown"))
        if label not in buckets:
            buckets[label] = AttributionBucket(label=label)

        b = buckets[label]
        pos_pnl = sum(pos.daily_pnl)
        pos_bps = pos_pnl / nav * 10_000
        b.total_pnl_usd += pos_pnl
        b.total_pnl_bps += pos_bps
        b.position_count += 1
        if pos_pnl > 0:
            b.win_count += 1
        else:
            b.loss_count += 1

    # Compute averages and contribution %
    for b in buckets.values():
        b.avg_return_bps = round(b.total_pnl_bps / max(b.position_count, 1), 2)
        b.contribution_pct = round(
            b.total_pnl_usd / abs(total_pnl) * 100, 1
        ) if total_pnl != 0 else 0.0
        b.total_pnl_usd = round(b.total_pnl_usd, 0)
        b.total_pnl_bps = round(b.total_pnl_bps, 2)

    return sorted(buckets.values(), key=lambda x: -abs(x.total_pnl_usd))


def compute_conviction_win_rate(positions: list[SimPosition]) -> dict:
    """Win rate broken down by conviction level."""
    by_conv: dict[int, list] = {}
    for pos in positions:
        c = pos.conviction
        if c not in by_conv:
            by_conv[c] = []
        pos_pnl = sum(pos.daily_pnl)
        by_conv[c].append(pos_pnl)

    result = {}
    for c in sorted(by_conv.keys()):
        pnls = by_conv[c]
        wins = sum(1 for p in pnls if p > 0)
        result[c] = {
            "conviction": c,
            "count": len(pnls),
            "wins": wins,
            "losses": len(pnls) - wins,
            "win_rate_pct": round(wins / max(len(pnls), 1) * 100, 1),
            "avg_pnl_usd": round(mean(pnls), 0),
            "total_pnl_usd": round(sum(pnls), 0),
        }
    return result


def compute_monthly_returns(
    positions: list[SimPosition],
    n_days: int,
    start_date: datetime,
) -> list[dict]:
    """Break daily P&L into monthly buckets."""

    nav = NAV_MILLIONS * 1_000_000

    # Aggregate daily P&L
    daily_pnl = [0.0] * n_days
    for pos in positions:
        for d in range(n_days):
            daily_pnl[d] += pos.daily_pnl[d]

    # Group by month
    months: dict[str, dict] = {}
    for d in range(n_days):
        dt = start_date + timedelta(days=d * 7 // 5)  # approx business days
        month_key = dt.strftime("%Y-%m")
        if month_key not in months:
            months[month_key] = {
                "month": month_key,
                "pnl_usd": 0.0,
                "pnl_bps": 0.0,
                "trading_days": 0,
                "winning_days": 0,
                "losing_days": 0,
            }
        m = months[month_key]
        m["pnl_usd"] += daily_pnl[d]
        m["pnl_bps"] += daily_pnl[d] / nav * 10_000
        m["trading_days"] += 1
        if daily_pnl[d] > 0:
            m["winning_days"] += 1
        else:
            m["losing_days"] += 1

    result = []
    for m in months.values():
        m["pnl_usd"] = round(m["pnl_usd"], 0)
        m["pnl_bps"] = round(m["pnl_bps"], 2)
        m["return_pct"] = round(m["pnl_bps"] / 100, 4)
        result.append(m)

    return result


def compute_cumulative_pnl(
    positions: list[SimPosition],
    n_days: int,
) -> list[float]:
    """Return cumulative P&L in USD for each day."""
    daily_pnl = [0.0] * n_days
    for pos in positions:
        for d in range(n_days):
            daily_pnl[d] += pos.daily_pnl[d]

    cum = []
    running = 0.0
    for p in daily_pnl:
        running += p
        cum.append(round(running, 0))
    return cum


# ---------------------------------------------------------------------------
# Excel output (5 sheets)
# ---------------------------------------------------------------------------

def _style_header_row(ws, headers, row=1):
    """Apply standard header styling."""
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = f"A{row + 1}"


def _auto_width(ws, min_width=10, max_width=35):
    """Auto-fit column widths."""
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[letter].width = min(max(max_len + 3, min_width), max_width)


def write_excel(
    metrics: BacktestMetrics,
    positions: list[SimPosition],
    monthly: list[dict],
    cum_pnl: list[float],
    sector_attr: list[AttributionBucket],
    rating_attr: list[AttributionBucket],
    alpha_attr: list[AttributionBucket],
    conv_win: dict,
    n_days: int,
    filepath: str,
) -> None:
    """Write multi-sheet Excel workbook."""

    wb = Workbook()

    # ── Sheet 1: Summary ────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"

    ws.cell(row=1, column=1, value="Strategies in Credit — Backtest Summary").font = SECTION_FONT
    ws.cell(row=2, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    ws.cell(row=3, column=1, value=f"Simulation: {n_days} trading days | NAV: ${NAV_MILLIONS:.0f}M")

    # Key metrics table
    metric_data = [
        ("Total Return (bps)", f"{metrics.total_return_bps:+.1f}"),
        ("Total Return (%)", f"{metrics.total_return_pct:+.4f}%"),
        ("Annualised Return (%)", f"{metrics.annualised_return_pct:+.2f}%"),
        ("Sharpe Ratio", f"{metrics.sharpe_ratio:.2f}"),
        ("Sortino Ratio", f"{metrics.sortino_ratio:.2f}"),
        ("Calmar Ratio", f"{metrics.calmar_ratio:.2f}"),
        ("Max Drawdown (%)", f"{metrics.max_drawdown_pct:.4f}%"),
        ("Max Drawdown (bps)", f"{metrics.max_drawdown_bps:.1f}"),
        ("Win Rate", f"{metrics.win_rate_pct:.1f}%"),
        ("Profit Factor", f"{metrics.profit_factor:.2f}"),
        ("Total P&L ($)", f"${metrics.total_pnl_usd:,.0f}"),
        ("Best Day (bps)", f"{metrics.best_day_bps:+.2f}"),
        ("Worst Day (bps)", f"{metrics.worst_day_bps:+.2f}"),
        ("Daily Vol (bps)", f"{metrics.daily_vol_bps:.2f}"),
        ("Winning Days", f"{metrics.winning_days}"),
        ("Losing Days", f"{metrics.losing_days}"),
    ]

    for i, (label, value) in enumerate(metric_data, start=5):
        ws.cell(row=i, column=1, value=label).font = LABEL_FONT
        cell = ws.cell(row=i, column=2, value=value)
        # Colour code P&L values
        if "+" in str(value) and ("Return" in label or "P&L" in label or "Best" in label):
            cell.font = GREEN_FONT
        elif "-" in str(value) and ("Return" in label or "P&L" in label or "Worst" in label or "Drawdown" in label):
            cell.font = RED_FONT

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 20

    # ── Sheet 2: Monthly Returns ────────────────────────────────────────
    ws2 = wb.create_sheet("Monthly Returns")
    headers = ["Month", "P&L ($)", "Return (bps)", "Return (%)",
               "Trading Days", "Winning Days", "Losing Days"]
    _style_header_row(ws2, headers)

    for i, m in enumerate(monthly, start=2):
        ws2.cell(row=i, column=1, value=m["month"])
        ws2.cell(row=i, column=2, value=m["pnl_usd"]).number_format = NUMBER_FMT_0DP
        ws2.cell(row=i, column=3, value=m["pnl_bps"]).number_format = NUMBER_FMT_2DP
        ws2.cell(row=i, column=4, value=m["return_pct"]).number_format = "0.0000%"
        ws2.cell(row=i, column=5, value=m["trading_days"])
        ws2.cell(row=i, column=6, value=m["winning_days"])
        ws2.cell(row=i, column=7, value=m["losing_days"])

        # Colour code
        pnl_cell = ws2.cell(row=i, column=2)
        if m["pnl_usd"] > 0:
            pnl_cell.font = GREEN_FONT
        elif m["pnl_usd"] < 0:
            pnl_cell.font = RED_FONT

    _auto_width(ws2)

    # ── Sheet 3: Position Detail ────────────────────────────────────────
    ws3 = wb.create_sheet("Position Detail")
    headers = ["Entity", "Direction", "Conv", "Notional ($M)",
               "Entry Spread", "Exit Spread", "Spread Δ",
               "P&L ($)", "P&L (bps)", "Alpha Source", "Sector", "Rating",
               "Outcome"]
    _style_header_row(ws3, headers)

    nav = NAV_MILLIONS * 1_000_000

    for i, pos in enumerate(sorted(positions, key=lambda p: -abs(sum(p.daily_pnl))), start=2):
        pos_pnl = sum(pos.daily_pnl)
        pos_bps = pos_pnl / nav * 10_000
        exit_spread = pos.spread_path[-1] if pos.spread_path else pos.entry_spread
        spread_delta = exit_spread - pos.entry_spread

        ws3.cell(row=i, column=1, value=pos.entity_name)
        dir_cell = ws3.cell(row=i, column=2, value=pos.direction)
        if pos.direction == "LONG_RISK":
            dir_cell.fill = LONG_FILL
        else:
            dir_cell.fill = SHORT_FILL
        ws3.cell(row=i, column=3, value=pos.conviction)
        ws3.cell(row=i, column=4, value=pos.notional_m).number_format = NUMBER_FMT_2DP
        ws3.cell(row=i, column=5, value=round(pos.entry_spread, 1))
        ws3.cell(row=i, column=6, value=round(exit_spread, 1))
        ws3.cell(row=i, column=7, value=round(spread_delta, 1))
        pnl_cell = ws3.cell(row=i, column=8, value=round(pos_pnl, 0))
        pnl_cell.number_format = NUMBER_FMT_0DP
        pnl_cell.font = GREEN_FONT if pos_pnl > 0 else RED_FONT
        ws3.cell(row=i, column=9, value=round(pos_bps, 2)).number_format = NUMBER_FMT_2DP
        ws3.cell(row=i, column=10, value=pos.alpha_source)
        ws3.cell(row=i, column=11, value=pos.sector)
        ws3.cell(row=i, column=12, value=pos.estimated_rating)
        outcome = "WIN" if pos_pnl > 0 else "LOSS"
        ws3.cell(row=i, column=13, value=outcome).font = GREEN_FONT if outcome == "WIN" else RED_FONT

    _auto_width(ws3)

    # ── Sheet 4: Attribution ────────────────────────────────────────────
    ws4 = wb.create_sheet("Attribution")

    row = 1
    ws4.cell(row=row, column=1, value="BY ALPHA SOURCE").font = SECTION_FONT
    row += 1
    attr_headers = ["Group", "P&L ($)", "P&L (bps)", "Positions",
                    "Wins", "Losses", "Avg Return (bps)", "Contribution (%)"]
    _style_header_row(ws4, attr_headers, row=row)
    row += 1

    for a in alpha_attr:
        ws4.cell(row=row, column=1, value=a.label)
        ws4.cell(row=row, column=2, value=a.total_pnl_usd).number_format = NUMBER_FMT_0DP
        ws4.cell(row=row, column=3, value=a.total_pnl_bps).number_format = NUMBER_FMT_2DP
        ws4.cell(row=row, column=4, value=a.position_count)
        ws4.cell(row=row, column=5, value=a.win_count)
        ws4.cell(row=row, column=6, value=a.loss_count)
        ws4.cell(row=row, column=7, value=a.avg_return_bps).number_format = NUMBER_FMT_2DP
        ws4.cell(row=row, column=8, value=a.contribution_pct).number_format = NUMBER_FMT_2DP
        row += 1

    row += 2
    ws4.cell(row=row, column=1, value="BY SECTOR").font = SECTION_FONT
    row += 1
    _style_header_row(ws4, attr_headers, row=row)
    row += 1
    for a in sector_attr:
        ws4.cell(row=row, column=1, value=a.label)
        ws4.cell(row=row, column=2, value=a.total_pnl_usd).number_format = NUMBER_FMT_0DP
        ws4.cell(row=row, column=3, value=a.total_pnl_bps).number_format = NUMBER_FMT_2DP
        ws4.cell(row=row, column=4, value=a.position_count)
        ws4.cell(row=row, column=5, value=a.win_count)
        ws4.cell(row=row, column=6, value=a.loss_count)
        ws4.cell(row=row, column=7, value=a.avg_return_bps).number_format = NUMBER_FMT_2DP
        ws4.cell(row=row, column=8, value=a.contribution_pct).number_format = NUMBER_FMT_2DP
        row += 1

    row += 2
    ws4.cell(row=row, column=1, value="BY RATING BUCKET").font = SECTION_FONT
    row += 1
    _style_header_row(ws4, attr_headers, row=row)
    row += 1
    for a in rating_attr:
        ws4.cell(row=row, column=1, value=a.label)
        ws4.cell(row=row, column=2, value=a.total_pnl_usd).number_format = NUMBER_FMT_0DP
        ws4.cell(row=row, column=3, value=a.total_pnl_bps).number_format = NUMBER_FMT_2DP
        ws4.cell(row=row, column=4, value=a.position_count)
        ws4.cell(row=row, column=5, value=a.win_count)
        ws4.cell(row=row, column=6, value=a.loss_count)
        ws4.cell(row=row, column=7, value=a.avg_return_bps).number_format = NUMBER_FMT_2DP
        ws4.cell(row=row, column=8, value=a.contribution_pct).number_format = NUMBER_FMT_2DP
        row += 1

    _auto_width(ws4)

    # ── Sheet 5: Conviction Analysis ────────────────────────────────────
    ws5 = wb.create_sheet("Conviction Analysis")
    headers = ["Conviction", "Positions", "Wins", "Losses",
               "Win Rate (%)", "Avg P&L ($)", "Total P&L ($)"]
    _style_header_row(ws5, headers)

    for i, (c, data) in enumerate(sorted(conv_win.items()), start=2):
        ws5.cell(row=i, column=1, value=f"⚫ × {c}")
        ws5.cell(row=i, column=2, value=data["count"])
        ws5.cell(row=i, column=3, value=data["wins"])
        ws5.cell(row=i, column=4, value=data["losses"])
        ws5.cell(row=i, column=5, value=data["win_rate_pct"])
        ws5.cell(row=i, column=6, value=data["avg_pnl_usd"]).number_format = NUMBER_FMT_0DP
        ws5.cell(row=i, column=7, value=data["total_pnl_usd"]).number_format = NUMBER_FMT_0DP

    _auto_width(ws5)

    wb.save(filepath)


# ---------------------------------------------------------------------------
# PDF output (one-page summary)
# ---------------------------------------------------------------------------

def get_pdf_styles() -> dict:
    """Build paragraph styles for the PDF."""
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "title", parent=base["Normal"],
            fontName="Helvetica-Bold", fontSize=16,
            textColor=colors.white, leading=20,
        ),
        "subtitle": ParagraphStyle(
            "subtitle", parent=base["Normal"],
            fontName="Helvetica", fontSize=9,
            textColor=LIGHT_STEEL, leading=12,
        ),
        "section": ParagraphStyle(
            "section", parent=base["Normal"],
            fontName="Helvetica-Bold", fontSize=10,
            textColor=NAVY, leading=13,
            spaceBefore=6, spaceAfter=3,
        ),
        "body": ParagraphStyle(
            "body", parent=base["Normal"],
            fontName="Helvetica", fontSize=8.5,
            textColor=colors.HexColor("#212529"), leading=11,
        ),
        "metric_label": ParagraphStyle(
            "metric_label", parent=base["Normal"],
            fontName="Helvetica", fontSize=7.5,
            textColor=STEEL, leading=10, alignment=TA_CENTER,
        ),
        "metric_value": ParagraphStyle(
            "metric_value", parent=base["Normal"],
            fontName="Helvetica-Bold", fontSize=13,
            textColor=NAVY, leading=16, alignment=TA_CENTER,
        ),
        "footer": ParagraphStyle(
            "footer", parent=base["Normal"],
            fontName="Helvetica", fontSize=6.5,
            textColor=LIGHT_STEEL, alignment=TA_CENTER,
        ),
    }


def _draw_header(canvas, doc, n_days):
    """Draw the dark header bar."""
    width, height = A4
    header_h = 50
    canvas.setFillColor(DARK_NAVY)
    canvas.rect(0, height - header_h, width, header_h, fill=1, stroke=0)

    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 16)
    canvas.drawString(1.5 * cm, height - 24, "Strategies in Credit")

    canvas.setFillColor(LIGHT_STEEL)
    canvas.setFont("Helvetica", 9)
    date_str = datetime.now().strftime("%d %B %Y")
    canvas.drawString(1.5 * cm, height - 40,
                      f"Simulated Track Record  |  {n_days} Trading Days  |  {date_str}")

    # Badge
    canvas.setFillColor(ACCENT)
    badge_w, badge_h = 80, 20
    badge_x = width - badge_w - 1.5 * cm
    badge_y = height - 38
    canvas.roundRect(badge_x, badge_y, badge_w, badge_h, 3, fill=1, stroke=0)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawCentredString(badge_x + badge_w / 2, badge_y + 6, "BACKTEST")


def _draw_footer(canvas, doc):
    """Draw the footer."""
    width, _ = A4
    canvas.setFillColor(LIGHT_STEEL)
    canvas.setFont("Helvetica", 6.5)
    canvas.drawCentredString(
        width / 2, 15,
        "Strategies in Credit  |  Simulated results do not represent actual trading  |  Past simulated performance is not indicative of future results"
    )


def _metric_cell(value, label, styles):
    """Create a metric cell (value on top, label below)."""
    return Table(
        [
            [Paragraph(str(value), styles["metric_value"])],
            [Paragraph(label, styles["metric_label"])],
        ],
        colWidths=[3.8 * cm],
        rowHeights=[0.55 * cm, 0.35 * cm],
        style=TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 1),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]),
    )


def write_pdf(
    metrics: BacktestMetrics,
    positions: list[SimPosition],
    monthly: list[dict],
    sector_attr: list[AttributionBucket],
    alpha_attr: list[AttributionBucket],
    conv_win: dict,
    n_days: int,
    filepath: str,
) -> None:
    """Write one-page PDF summary."""

    styles = get_pdf_styles()
    width, height = A4

    # Create document with header/footer
    doc = BaseDocTemplate(
        filepath,
        pagesize=A4,
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=55,
        bottomMargin=30,
    )

    content_width = width - 2.4 * cm
    frame = Frame(
        doc.leftMargin, doc.bottomMargin,
        content_width, height - 55 - 30,
        id="main",
    )
    template = PageTemplate(
        "main", frames=[frame],
        onPage=lambda c, d: (_draw_header(c, d, n_days), _draw_footer(c, d)),
    )
    doc.addPageTemplates([template])

    story = []

    # ── Key Metrics Row ─────────────────────────────────────────────────
    sharpe_color = LONG_GREEN if metrics.sharpe_ratio > 1.0 else (SHORT_RED if metrics.sharpe_ratio < 0 else NAVY)

    row1 = Table(
        [[
            _metric_cell(f"{metrics.total_return_bps:+.0f}", "Total Return (bps)", styles),
            _metric_cell(f"{metrics.sharpe_ratio:.2f}", "Sharpe Ratio", styles),
            _metric_cell(f"{metrics.sortino_ratio:.2f}", "Sortino Ratio", styles),
            _metric_cell(f"{metrics.max_drawdown_pct:.2f}%", "Max Drawdown", styles),
            _metric_cell(f"{metrics.win_rate_pct:.0f}%", "Win Rate", styles),
        ]],
        colWidths=[content_width / 5] * 5,
        style=TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOX", (0, 0), (-1, -1), 0.5, DIVIDER),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, DIVIDER),
            ("BACKGROUND", (0, 0), (-1, -1), OFF_WHITE),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]),
    )
    story.append(row1)
    story.append(Spacer(1, 6))

    # ── Secondary Metrics Row ───────────────────────────────────────────
    row2 = Table(
        [[
            _metric_cell(f"${metrics.total_pnl_usd:,.0f}", "Total P&L", styles),
            _metric_cell(f"{metrics.annualised_return_pct:+.1f}%", "Ann. Return", styles),
            _metric_cell(f"{metrics.profit_factor:.2f}", "Profit Factor", styles),
            _metric_cell(f"{metrics.calmar_ratio:.2f}", "Calmar Ratio", styles),
            _metric_cell(f"{metrics.daily_vol_bps:.1f}bp", "Daily Vol", styles),
        ]],
        colWidths=[content_width / 5] * 5,
        style=TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOX", (0, 0), (-1, -1), 0.5, DIVIDER),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, DIVIDER),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]),
    )
    story.append(row2)
    story.append(Spacer(1, 8))

    # ── Monthly Returns Table ───────────────────────────────────────────
    story.append(Paragraph("MONTHLY RETURNS", styles["section"]))

    monthly_data = [["Month", "P&L ($k)", "Return (bps)", "Win Days", "Loss Days"]]
    for m in monthly:
        pnl_k = f"${m['pnl_usd'] / 1000:,.0f}k"
        monthly_data.append([
            m["month"],
            pnl_k,
            f"{m['pnl_bps']:+.1f}",
            str(m["winning_days"]),
            str(m["losing_days"]),
        ])

    monthly_table = Table(
        monthly_data,
        colWidths=[3 * cm, 3 * cm, 3 * cm, 2.5 * cm, 2.5 * cm],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, DIVIDER),
            ("ROWHEIGHT", (0, 0), (-1, -1), 14),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]),
    )
    story.append(monthly_table)
    story.append(Spacer(1, 8))

    # ── Attribution by Alpha Source ──────────────────────────────────────
    story.append(Paragraph("P&L ATTRIBUTION BY ALPHA SOURCE", styles["section"]))

    alpha_data = [["Source", "P&L ($k)", "P&L (bps)", "Positions", "Win Rate", "Contribution"]]
    for a in alpha_attr:
        wr = round(a.win_count / max(a.position_count, 1) * 100, 0)
        alpha_data.append([
            a.label,
            f"${a.total_pnl_usd / 1000:,.0f}k",
            f"{a.total_pnl_bps:+.1f}",
            str(a.position_count),
            f"{wr:.0f}%",
            f"{a.contribution_pct:+.1f}%",
        ])

    alpha_table = Table(
        alpha_data,
        colWidths=[3.2 * cm, 2.5 * cm, 2.5 * cm, 2 * cm, 2 * cm, 2.5 * cm],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, DIVIDER),
            ("ROWHEIGHT", (0, 0), (-1, -1), 14),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]),
    )
    story.append(alpha_table)
    story.append(Spacer(1, 8))

    # ── Attribution by Sector ───────────────────────────────────────────
    story.append(Paragraph("P&L ATTRIBUTION BY SECTOR", styles["section"]))

    sector_data = [["Sector", "P&L ($k)", "P&L (bps)", "Positions", "Win Rate", "Contribution"]]
    for a in sector_attr:
        wr = round(a.win_count / max(a.position_count, 1) * 100, 0)
        sector_data.append([
            a.label,
            f"${a.total_pnl_usd / 1000:,.0f}k",
            f"{a.total_pnl_bps:+.1f}",
            str(a.position_count),
            f"{wr:.0f}%",
            f"{a.contribution_pct:+.1f}%",
        ])

    sector_table = Table(
        sector_data,
        colWidths=[3.2 * cm, 2.5 * cm, 2.5 * cm, 2 * cm, 2 * cm, 2.5 * cm],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, DIVIDER),
            ("ROWHEIGHT", (0, 0), (-1, -1), 14),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]),
    )
    story.append(sector_table)
    story.append(Spacer(1, 8))

    # ── Conviction Win Rate ─────────────────────────────────────────────
    story.append(Paragraph("WIN RATE BY CONVICTION LEVEL", styles["section"]))

    conv_data = [["Conviction", "Positions", "Wins", "Losses", "Win Rate", "Total P&L ($k)"]]
    for c in sorted(conv_win.keys()):
        d = conv_win[c]
        conv_data.append([
            f"{'*' * c} ({c}/5)",
            str(d["count"]),
            str(d["wins"]),
            str(d["losses"]),
            f"{d['win_rate_pct']:.0f}%",
            f"${d['total_pnl_usd'] / 1000:,.0f}k",
        ])

    conv_table = Table(
        conv_data,
        colWidths=[3.2 * cm, 2 * cm, 2 * cm, 2 * cm, 2 * cm, 3 * cm],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, DIVIDER),
            ("ROWHEIGHT", (0, 0), (-1, -1), 14),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]),
    )
    story.append(conv_table)
    story.append(Spacer(1, 8))

    # ── Top/Bottom Positions ────────────────────────────────────────────
    story.append(Paragraph("TOP & BOTTOM POSITIONS", styles["section"]))

    nav = NAV_MILLIONS * 1_000_000
    sorted_pos = sorted(positions, key=lambda p: sum(p.daily_pnl), reverse=True)
    top3 = sorted_pos[:3]
    bot3 = sorted_pos[-3:]

    pos_data = [["Entity", "Direction", "Entry", "Exit", "P&L ($k)", "P&L (bps)"]]
    for pos in top3 + bot3:
        pos_pnl = sum(pos.daily_pnl)
        exit_s = pos.spread_path[-1] if pos.spread_path else pos.entry_spread
        pos_data.append([
            pos.entity_name[:25],
            pos.direction.replace("_RISK", ""),
            f"{pos.entry_spread:.0f}",
            f"{exit_s:.0f}",
            f"${pos_pnl / 1000:,.0f}k",
            f"{pos_pnl / nav * 10000:+.1f}",
        ])

    pos_table = Table(
        pos_data,
        colWidths=[4.5 * cm, 2 * cm, 1.5 * cm, 1.5 * cm, 2.5 * cm, 2.5 * cm],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, DIVIDER),
            ("ROWHEIGHT", (0, 0), (-1, -1), 14),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            # Green top 3
            ("BACKGROUND", (0, 1), (-1, 3), colors.HexColor("#E8F5E9")),
            # Red bottom 3
            ("BACKGROUND", (0, 4), (-1, 6), colors.HexColor("#FFEBEE")),
        ]),
    )
    story.append(pos_table)

    # ── Disclaimer ──────────────────────────────────────────────────────
    story.append(Spacer(1, 10))
    disclaimer = (
        "<i>This simulated track record assumes 50% mean-reversion of spreads toward "
        "analyst fair values over the simulation horizon with 5bp daily volatility noise. "
        "It does not represent actual trading results. Transaction costs, slippage, "
        "and market impact are not modelled. Past simulated performance is not "
        "indicative of future results.</i>"
    )
    story.append(Paragraph(disclaimer, ParagraphStyle(
        "disclaimer", fontName="Helvetica", fontSize=6.5,
        textColor=STEEL, leading=9, alignment=TA_CENTER,
    )))

    doc.build(story)


# ---------------------------------------------------------------------------
# Terminal output
# ---------------------------------------------------------------------------

def print_report(
    metrics: BacktestMetrics,
    positions: list[SimPosition],
    monthly: list[dict],
    sector_attr: list[AttributionBucket],
    alpha_attr: list[AttributionBucket],
    conv_win: dict,
    n_days: int,
) -> None:
    """Pretty-print backtest results to terminal."""

    nav = NAV_MILLIONS * 1_000_000
    W = 72

    print()
    print("=" * W)
    print("  STRATEGIES IN CREDIT — SIMULATED TRACK RECORD")
    print(f"  {n_days} trading days | NAV: ${NAV_MILLIONS:.0f}M | "
          f"{datetime.now().strftime('%Y-%m-%d')}")
    print("=" * W)

    # Key metrics
    print()
    print("  KEY METRICS")
    print("  " + "-" * 40)
    pnl_sign = "+" if metrics.total_pnl_usd > 0 else ""
    print(f"  Total Return:      {metrics.total_return_bps:+.1f} bps "
          f"({metrics.total_return_pct:+.4f}%)")
    print(f"  Total P&L:         {pnl_sign}${metrics.total_pnl_usd:,.0f}")
    print(f"  Annualised Return: {metrics.annualised_return_pct:+.2f}%")
    print(f"  Sharpe Ratio:      {metrics.sharpe_ratio:.2f}")
    print(f"  Sortino Ratio:     {metrics.sortino_ratio:.2f}")
    print(f"  Calmar Ratio:      {metrics.calmar_ratio:.2f}")
    print(f"  Max Drawdown:      {metrics.max_drawdown_pct:.4f}% "
          f"({metrics.max_drawdown_bps:.1f} bps)")
    print(f"  Win Rate:          {metrics.win_rate_pct:.1f}% "
          f"({metrics.winning_days}W / {metrics.losing_days}L)")
    print(f"  Profit Factor:     {metrics.profit_factor:.2f}")
    print(f"  Daily Vol:         {metrics.daily_vol_bps:.2f} bps")
    print(f"  Best Day:          {metrics.best_day_bps:+.2f} bps")
    print(f"  Worst Day:         {metrics.worst_day_bps:+.2f} bps")

    # Monthly returns
    print()
    print("  MONTHLY RETURNS")
    print("  " + "-" * 55)
    print(f"  {'Month':<10} {'P&L ($k)':>12} {'Return (bps)':>14} {'Win/Loss':>10}")
    for m in monthly:
        pnl_k = m["pnl_usd"] / 1000
        print(f"  {m['month']:<10} {pnl_k:>+11,.0f}k {m['pnl_bps']:>+13.1f} "
              f"{m['winning_days']:>4}W/{m['losing_days']}L")

    # Attribution by alpha source
    print()
    print("  P&L ATTRIBUTION BY ALPHA SOURCE")
    print("  " + "-" * 55)
    print(f"  {'Source':<16} {'P&L ($k)':>10} {'P&L (bps)':>10} {'Pos':>5} {'Contr%':>8}")
    for a in alpha_attr:
        pnl_k = a.total_pnl_usd / 1000
        print(f"  {a.label:<16} {pnl_k:>+9,.0f}k {a.total_pnl_bps:>+9.1f} "
              f"{a.position_count:>5} {a.contribution_pct:>+7.1f}%")

    # Attribution by sector
    print()
    print("  P&L ATTRIBUTION BY SECTOR")
    print("  " + "-" * 55)
    print(f"  {'Sector':<22} {'P&L ($k)':>10} {'P&L (bps)':>10} {'Pos':>5} {'Contr%':>8}")
    for a in sector_attr:
        pnl_k = a.total_pnl_usd / 1000
        print(f"  {a.label:<22} {pnl_k:>+9,.0f}k {a.total_pnl_bps:>+9.1f} "
              f"{a.position_count:>5} {a.contribution_pct:>+7.1f}%")

    # Conviction win rate
    print()
    print("  WIN RATE BY CONVICTION")
    print("  " + "-" * 55)
    print(f"  {'Conv':>5} {'Pos':>5} {'W':>4} {'L':>4} {'Win%':>7} {'Total P&L ($k)':>16}")
    for c in sorted(conv_win.keys()):
        d = conv_win[c]
        pnl_k = d["total_pnl_usd"] / 1000
        dots = "*" * c
        print(f"  {dots:>5} {d['count']:>5} {d['wins']:>4} {d['losses']:>4} "
              f"{d['win_rate_pct']:>6.0f}% {pnl_k:>+15,.0f}k")

    # Top/bottom positions
    print()
    print("  TOP WINNERS & LOSERS")
    print("  " + "-" * 55)
    sorted_pos = sorted(positions, key=lambda p: sum(p.daily_pnl), reverse=True)

    for tag, group in [("TOP 3", sorted_pos[:3]), ("BOTTOM 3", sorted_pos[-3:])]:
        print(f"  {tag}:")
        for pos in group:
            pos_pnl = sum(pos.daily_pnl)
            pos_bps = pos_pnl / nav * 10_000
            exit_s = pos.spread_path[-1] if pos.spread_path else pos.entry_spread
            dir_tag = "L" if pos.direction == "LONG_RISK" else "S"
            print(f"    {dir_tag} {pos.entity_name:<25} "
                  f"{pos.entry_spread:>6.0f} -> {exit_s:>6.0f}  "
                  f"${pos_pnl / 1000:>+8,.0f}k  ({pos_bps:>+5.1f}bp)")

    print()
    print("=" * W)


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Strategies in Credit — Simulated Track Record Generator"
    )
    parser.add_argument("--days", type=int, default=DEFAULT_TRADING_DAYS,
                        help=f"Trading days to simulate (default: {DEFAULT_TRADING_DAYS})")
    parser.add_argument("--seed", type=int, default=None,
                        help="Random seed for reproducible results")
    parser.add_argument("--json", action="store_true",
                        help="Output raw JSON instead of formatted report")
    parser.add_argument("--no-excel", action="store_true",
                        help="Skip Excel output")
    parser.add_argument("--no-pdf", action="store_true",
                        help="Skip PDF output")
    parser.add_argument("--historical", action="store_true",
                        help="Use historical spread data (requires data/spread_history/)")
    args = parser.parse_args()

    # Historical mode placeholder
    if args.historical:
        hist_dir = Path("data/spread_history")
        if not hist_dir.exists():
            print("ERROR: --historical requires data/spread_history/ directory")
            print("       Upload daily spread snapshots as CSV/Excel files to this directory.")
            print("       Format: data/spread_history/spreads_YYYYMMDD.csv")
            print("       Columns: entity_name, spread_bps, date")
            print()
            print("       Once available, this will run a proper walk-forward backtest")
            print("       using actual daily spread movements instead of simulated paths.")
            sys.exit(1)
        print("Historical backtest mode not yet implemented.")
        print("Upload spread history data and this feature will be enabled.")
        sys.exit(0)

    # Load portfolio
    portfolio_path = find_latest_portfolio()
    if not portfolio_path:
        print("ERROR: No portfolio found in outputs/portfolio/")
        print("       Run: python -m agents.strategist")
        sys.exit(1)

    portfolio = load_portfolio(portfolio_path)
    positions = build_positions(portfolio)

    if not positions:
        print("ERROR: No positions found in portfolio")
        sys.exit(1)

    print(f"Loaded {len(positions)} positions from {os.path.basename(portfolio_path)}")

    # Run simulation
    n_days = args.days
    rng = np.random.default_rng(args.seed)

    print(f"Simulating {n_days} trading days "
          f"(seed={'random' if args.seed is None else args.seed})...")

    simulate_spread_paths(positions, n_days, rng)

    # Compute all metrics
    metrics = compute_metrics(positions, n_days)
    total_pnl = metrics.total_pnl_usd

    monthly = compute_monthly_returns(positions, n_days, datetime.now())
    cum_pnl = compute_cumulative_pnl(positions, n_days)

    sector_attr = compute_attribution(positions, "sector", total_pnl)
    rating_attr = compute_attribution(positions, "estimated_rating", total_pnl)
    alpha_attr = compute_attribution(positions, "alpha_source", total_pnl)
    conv_win = compute_conviction_win_rate(positions)

    # JSON output
    if args.json:
        output = {
            "metrics": asdict(metrics),
            "monthly_returns": monthly,
            "cumulative_pnl": cum_pnl,
            "attribution_by_source": [asdict(a) for a in alpha_attr],
            "attribution_by_sector": [asdict(a) for a in sector_attr],
            "attribution_by_rating": [asdict(a) for a in rating_attr],
            "conviction_analysis": conv_win,
            "positions": [
                {
                    "entity_name": p.entity_name,
                    "direction": p.direction,
                    "conviction": p.conviction,
                    "entry_spread": p.entry_spread,
                    "exit_spread": p.spread_path[-1] if p.spread_path else p.entry_spread,
                    "pnl_usd": round(sum(p.daily_pnl), 0),
                    "alpha_source": p.alpha_source,
                    "sector": p.sector,
                }
                for p in positions
            ],
        }
        print(json.dumps(output, indent=2, default=str))
        return

    # Terminal report
    print_report(metrics, positions, monthly, sector_attr, alpha_attr, conv_win, n_days)

    # Excel output
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    date_str = datetime.now().strftime("%Y%m%d")

    if not args.no_excel:
        excel_path = str(OUTPUT_DIR / f"backtest_report_{date_str}.xlsx")
        write_excel(
            metrics, positions, monthly, cum_pnl,
            sector_attr, rating_attr, alpha_attr, conv_win,
            n_days, excel_path,
        )
        print(f"\n  Excel: {excel_path}")

    # PDF output
    if not args.no_pdf:
        pdf_path = str(OUTPUT_DIR / f"backtest_report_{date_str}.pdf")
        write_pdf(
            metrics, positions, monthly,
            sector_attr, alpha_attr, conv_win,
            n_days, pdf_path,
        )
        print(f"  PDF:   {pdf_path}")

    print()


if __name__ == "__main__":
    main()
